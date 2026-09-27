"""Generate a reproducible, entirely fictional KMB Manufacturing source workbook."""

from __future__ import annotations

import hashlib
import json
import random
from collections import defaultdict
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo


ROOT = Path(__file__).resolve().parents[1]
DATA_DIR = ROOT / "data"
WORKBOOK_PATH = DATA_DIR / "kmb_synthetic_operations.xlsx"
MANIFEST_PATH = DATA_DIR / "kmb_synthetic_operations_manifest.json"
SEED = 20260918
SOURCE_LABEL = "Synthetic KMB Demonstration"
START_DATE = date(2024, 9, 1)
END_DATE = date(2026, 8, 31)


def date_key(value: date) -> int:
    return int(value.strftime("%Y%m%d"))


def week_start(value: date) -> date:
    return value - timedelta(days=value.weekday())


def rows_to_sheet(workbook: Workbook, title: str, columns: list[str], rows: list[dict]) -> None:
    sheet = workbook.create_sheet(title)
    sheet.append(columns)
    for cell in sheet[1]:
        cell.font = Font(name="Arial", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    for row in rows:
        sheet.append([row.get(column) for column in columns])
    for column_cells in sheet.columns:
        width = min(max(len(str(cell.value or "")) for cell in column_cells) + 2, 32)
        sheet.column_dimensions[column_cells[0].column_letter].width = width
    sheet.freeze_panes = "A2"
    table = Table(displayName=title, ref=f"A1:{sheet.cell(sheet.max_row, sheet.max_column).coordinate}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
    sheet.add_table(table)


def create_dimensions() -> dict[str, list[dict]]:
    dates = []
    value = START_DATE
    while value <= END_DATE:
        ws = week_start(value)
        dates.append({
            "DateKey": date_key(value), "Date": value, "Year": value.year, "Quarter": f"Q{((value.month - 1) // 3) + 1}",
            "MonthNumber": value.month, "Month": value.strftime("%B"), "YearMonth": value.strftime("%Y-%m"),
            "WeekStartDate": ws, "WeekLabel": f"{ws:%Y-%m-%d}", "DayOfWeekNumber": value.isoweekday(),
            "DayName": value.strftime("%A"), "IsWeekend": value.weekday() >= 5,
            "IsCompleteWeek": ws >= START_DATE and ws + timedelta(days=6) <= END_DATE,
        })
        value += timedelta(days=1)

    products = [
        {"ProductKey": 1, "ProductCode": "FG-S100", "ProductFamily": "Standard", "ProductName": "FlowGuard Standard 100", "Revision": "R1", "StandardCycleMinutes": 18.0, "StandardUnitValue": 425.0, "StandardUnitCost": 245.0},
        {"ProductKey": 2, "ProductCode": "FG-S200", "ProductFamily": "Standard", "ProductName": "FlowGuard Standard 200", "Revision": "R1", "StandardCycleMinutes": 24.0, "StandardUnitValue": 610.0, "StandardUnitCost": 354.0},
        {"ProductKey": 3, "ProductCode": "FG-H300", "ProductFamily": "High-Flow", "ProductName": "FlowGuard High-Flow 300", "Revision": "R2", "StandardCycleMinutes": 31.0, "StandardUnitValue": 890.0, "StandardUnitCost": 516.0},
        {"ProductKey": 4, "ProductCode": "FG-H500", "ProductFamily": "High-Flow", "ProductName": "FlowGuard High-Flow 500", "Revision": "R2", "StandardCycleMinutes": 42.0, "StandardUnitValue": 1230.0, "StandardUnitCost": 713.0},
        {"ProductKey": 5, "ProductCode": "FG-C200", "ProductFamily": "Chemical-Resistant", "ProductName": "FlowGuard Chemical 200", "Revision": "R1", "StandardCycleMinutes": 36.0, "StandardUnitValue": 1040.0, "StandardUnitCost": 603.0},
        {"ProductKey": 6, "ProductCode": "FG-C400", "ProductFamily": "Chemical-Resistant", "ProductName": "FlowGuard Chemical 400", "Revision": "R1", "StandardCycleMinutes": 48.0, "StandardUnitValue": 1490.0, "StandardUnitCost": 864.0},
    ]
    work_centers = [
        {"WorkCenterKey": 1, "WorkCenterCode": "WC-100", "Line": "Line A", "Process": "Cutting", "Department": "Operations"},
        {"WorkCenterKey": 2, "WorkCenterCode": "WC-120", "Line": "Line A", "Process": "Assembly", "Department": "Operations"},
        {"WorkCenterKey": 3, "WorkCenterCode": "WC-140", "Line": "Line A", "Process": "Leak Test", "Department": "Quality"},
        {"WorkCenterKey": 4, "WorkCenterCode": "WC-200", "Line": "Line B", "Process": "Assembly", "Department": "Operations"},
        {"WorkCenterKey": 5, "WorkCenterCode": "WC-220", "Line": "Line B", "Process": "Final Inspection", "Department": "Quality"},
        {"WorkCenterKey": 6, "WorkCenterCode": "WC-240", "Line": "Line B", "Process": "Packout", "Department": "Operations"},
    ]
    defects = [
        {"DefectKey": 1, "DefectCode": "DF-LEAK", "DefectName": "Seal Leakage", "Category": "Functional", "Severity": "High"},
        {"DefectKey": 2, "DefectCode": "DF-FIT", "DefectName": "Component Misfit", "Category": "Assembly", "Severity": "Medium"},
        {"DefectKey": 3, "DefectCode": "DF-SURFACE", "DefectName": "Surface Damage", "Category": "Appearance", "Severity": "Low"},
        {"DefectKey": 4, "DefectCode": "DF-THREAD", "DefectName": "Thread Damage", "Category": "Material", "Severity": "High"},
        {"DefectKey": 5, "DefectCode": "DF-LABEL", "DefectName": "Label Mismatch", "Category": "Documentation", "Severity": "Low"},
        {"DefectKey": 6, "DefectCode": "DF-PRESSURE", "DefectName": "Pressure Test Failure", "Category": "Functional", "Severity": "Critical"},
    ]
    suppliers = [
        {"SupplierKey": 1, "SupplierCode": "SUP-ALP", "SupplierName": "Alpine Polymer Works", "Commodity": "Polymer Housing", "Status": "Approved"},
        {"SupplierKey": 2, "SupplierCode": "SUP-BRX", "SupplierName": "Borex Metalcraft", "Commodity": "Metal Fittings", "Status": "Approved"},
        {"SupplierKey": 3, "SupplierCode": "SUP-CSD", "SupplierName": "Cascade Media Labs", "Commodity": "Filter Media", "Status": "Conditional"},
        {"SupplierKey": 4, "SupplierCode": "SUP-DLT", "SupplierName": "Delta Seal Systems", "Commodity": "Seals", "Status": "Approved"},
    ]
    employees = []
    roles = [("Production Operator", "Operations"), ("Production Operator", "Operations"), ("Quality Technician", "Quality"), ("Maintenance Technician", "Operations"), ("Material Coordinator", "Supply Chain"), ("Training Coordinator", "HR")]
    for index in range(1, 61):
        role, department = roles[(index - 1) % len(roles)]
        employees.append({
            "EmployeeKey": index, "EmployeeCode": f"EMP-{index:03d}", "EmployeeName": f"KMB Team Member {index:03d}",
            "Role": role, "Department": department, "Shift": "Day" if index % 3 else "Evening",
            "HireDate": date(2020 + (index % 5), ((index * 3) % 12) + 1, ((index * 7) % 25) + 1), "IsActive": True,
        })
    return {"Dim_Date": dates, "Dim_Product": products, "Dim_WorkCenter": work_centers, "Dim_Employee": employees, "Dim_Defect": defects, "Dim_Supplier": suppliers}


def create_facts(dimensions: dict[str, list[dict]], rng: random.Random) -> dict[str, list[dict]]:
    dates = dimensions["Dim_Date"]
    products = dimensions["Dim_Product"]
    work_centers = dimensions["Dim_WorkCenter"]
    defects = dimensions["Dim_Defect"]
    suppliers = dimensions["Dim_Supplier"]
    employees = dimensions["Dim_Employee"]
    production, inspections, mrb_cases, actions, training, labor = [], [], [], [], [], []
    event_id = inspection_id = case_id = action_id = labor_id = 0
    product_by_key = {row["ProductKey"]: row for row in products}
    for current in dates:
        actual_date = current["Date"]
        if not current["IsWeekend"]:
            for wc in work_centers:
                for shift in ("Day", "Evening"):
                    event_id += 1
                    product = products[(event_id + wc["WorkCenterKey"] + actual_date.month) % len(products)]
                    seasonal = 1.06 if actual_date.month in (3, 4, 9, 10) else 0.94 if actual_date.month in (12, 1) else 1.0
                    planned = max(35, round((78 + rng.gauss(0, 9)) * seasonal))
                    produced = max(1, planned + rng.randint(-7, 5))
                    defect_rate = 0.009 + (0.005 if wc["Process"] in ("Leak Test", "Final Inspection") else 0) + (0.004 if actual_date.month in (1, 2) else 0)
                    scrap = min(produced, max(0, round(produced * max(0, rng.gauss(defect_rate, 0.004)))))
                    rework = min(produced - scrap, max(0, round(produced * max(0, rng.gauss(0.018, 0.007)))))
                    good = produced - scrap - rework
                    labor_hours = round(max(4, produced * product["StandardCycleMinutes"] / 60 * rng.uniform(0.96, 1.14)), 2)
                    downtime = max(0, round(rng.gauss(32 if wc["Process"] == "Assembly" else 20, 12)))
                    planned_completion = datetime.combine(actual_date, datetime.min.time()) + timedelta(hours=14 if shift == "Day" else 22)
                    # Deterministic timing preserves the established random sequence for every pre-existing field.
                    late_minutes = 45 + ((event_id * 11 + wc["WorkCenterKey"] * 7) % 91) if event_id % 20 == 0 else -1 - ((event_id * 13 + product["ProductKey"] * 5) % 75)
                    actual_completion = planned_completion + timedelta(minutes=late_minutes)
                    production.append({
                        "ProductionEventKey": event_id, "DateKey": current["DateKey"], "ProductKey": product["ProductKey"], "WorkCenterKey": wc["WorkCenterKey"], "Shift": shift,
                        "PlannedQty": planned, "ProducedQty": produced, "GoodQty": good, "ScrapQty": scrap, "ReworkQty": rework,
                        "LaborHours": labor_hours, "DowntimeMinutes": downtime, "ProductionValue": round(produced * product["StandardUnitValue"], 2),
                        "PlannedCompletionDateTime": planned_completion, "ActualCompletionDateTime": actual_completion, "DataSource": SOURCE_LABEL,
                    })
                    inspection_id += 1
                    defect = defects[(event_id * 3 + actual_date.day) % len(defects)]
                    defect_qty = scrap + max(0, round(rework * rng.uniform(0.25, 0.8)))
                    inspections.append({
                        "InspectionKey": inspection_id, "ProductionEventKey": event_id, "DateKey": current["DateKey"], "ProductKey": product["ProductKey"], "WorkCenterKey": wc["WorkCenterKey"],
                        "DefectKey": defect["DefectKey"], "InspectedQty": produced, "DefectQty": min(defect_qty, produced),
                        "InspectionResult": "Fail" if defect_qty else "Pass", "DataSource": SOURCE_LABEL,
                    })
                    if (defect_qty > 0 and rng.random() < 0.08) or (defect["Severity"] == "Critical" and defect_qty > 0 and rng.random() < 0.30):
                        case_id += 1
                        open_date = min(actual_date + timedelta(days=rng.randint(0, 2)), END_DATE)
                        is_open = rng.random() < 0.10 or open_date > END_DATE - timedelta(days=28)
                        close_date = None if is_open else open_date + timedelta(days=rng.randint(2, 46))
                        status = "Open" if is_open else "Closed"
                        supplier = suppliers[(case_id + defect["DefectKey"]) % len(suppliers)]
                        material_cost = round(defect_qty * product["StandardUnitValue"] * rng.uniform(0.25, 0.95), 2)
                        mrb_cases.append({
                            "MRBCaseKey": case_id, "DateKey": date_key(open_date), "ProductKey": product["ProductKey"], "WorkCenterKey": wc["WorkCenterKey"], "SupplierKey": supplier["SupplierKey"], "DefectKey": defect["DefectKey"],
                            "OpenDate": open_date, "CloseDate": close_date, "Quantity": defect_qty, "MaterialCost": material_cost,
                            "Disposition": rng.choice(["Rework", "Use As Is", "Return to Supplier", "Scrap"]), "Status": status, "DataSource": SOURCE_LABEL,
                        })
                        action_id += 1
                        due_date = open_date + timedelta(days=14)
                        action_closed = not is_open and rng.random() > 0.08
                        completion = (close_date or open_date) + timedelta(days=rng.randint(-4, 12)) if action_closed else None
                        actions.append({
                            "ActionKey": action_id, "MRBCaseKey": case_id, "OpenDate": open_date, "DueDate": due_date, "CompletionDate": completion,
                            "OwnerDepartment": wc["Department"], "Status": "Closed" if action_closed else "Open", "EffectivenessVerified": action_closed and rng.random() > 0.16, "DataSource": SOURCE_LABEL,
                        })
            for employee in employees:
                if employee["Department"] == "HR" and actual_date.weekday() > 3:
                    continue
                labor_id += 1
                is_absent = rng.random() < 0.025
                regular = 0 if is_absent else round(rng.uniform(7.2, 8.3), 2)
                overtime = 0 if is_absent else round(max(0, rng.gauss(0.45, 0.65)), 2)
                labor.append({
                    "LaborRecordKey": labor_id, "DateKey": current["DateKey"], "EmployeeKey": employee["EmployeeKey"],
                    "WorkCenterKey": ((employee["EmployeeKey"] - 1) % len(work_centers)) + 1, "RegularHours": regular, "OvertimeHours": overtime,
                    "AbsenceHours": 8.0 if is_absent else 0.0, "DataSource": SOURCE_LABEL,
                })
    training_id = 0
    courses = ["ISO 9001 Awareness", "Quality Inspection Fundamentals", "Safe Assembly Procedure", "Corrective Action Basics"]
    for employee in employees:
        for course_index, course in enumerate(courses):
            assignment = START_DATE + timedelta(days=(employee["EmployeeKey"] * 7 + course_index * 29) % 280)
            while assignment <= END_DATE:
                training_id += 1
                due = assignment + timedelta(days=30)
                completed = None if rng.random() < 0.06 else due + timedelta(days=rng.randint(-12, 18))
                status = "Open" if completed is None else "Completed"
                training.append({
                    "TrainingRecordKey": training_id, "EmployeeKey": employee["EmployeeKey"], "CourseName": course, "AssignmentDate": assignment,
                    "DueDate": due, "CompletionDate": completed, "Status": status, "ScorePct": None if completed is None else round(rng.uniform(76, 100), 1), "DataSource": SOURCE_LABEL,
                })
                assignment += timedelta(days=180)
    targets = []
    target_id = 0
    target_definitions = [
        ("On-Time Production %", 0.95, "HigherIsBetter"), ("First Pass Yield %", 0.98, "HigherIsBetter"), ("Defects per 1K Units", 12.0, "LowerIsBetter"),
        ("Scrap Rate %", 0.015, "LowerIsBetter"), ("Rework Rate %", 0.03, "LowerIsBetter"), ("MRB Open Cases", 15.0, "LowerIsBetter"),
        ("MRB Aging >30 Days", 0.0, "LowerIsBetter"), ("Cost of Poor Quality %", 0.022, "LowerIsBetter"), ("OEE Proxy %", 0.85, "HigherIsBetter"),
        ("Training On-Time %", 0.95, "HigherIsBetter"), ("Training Compliance %", 0.98, "HigherIsBetter"), ("Corrective Action On-Time %", 0.90, "HigherIsBetter"),
        ("Target Attainment %", 0.90, "HigherIsBetter"), ("Defect Recurrence Rate %", 0.20, "LowerIsBetter"),
    ]
    period = START_DATE.replace(day=1)
    while period <= END_DATE:
        for kpi_name, target_value, direction in target_definitions:
            target_id += 1
            targets.append({"TargetKey": target_id, "KPIName": kpi_name, "PeriodStart": period, "ScopeType": "Enterprise", "ScopeKey": "KMB", "TargetValue": target_value, "Direction": direction, "YellowTolerancePct": 0.05, "RedTolerancePct": 0.10, "IsActive": True, "DataSource": SOURCE_LABEL})
        period = (period.replace(day=28) + timedelta(days=4)).replace(day=1)
    return {"Fact_Production": production, "Fact_QualityInspection": inspections, "Fact_MRB": mrb_cases, "Fact_CorrectiveAction": actions, "Fact_Training": training, "Fact_Labor": labor, "Fact_KpiTargets": targets}


def generate() -> None:
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    rng = random.Random(SEED)
    dimensions = create_dimensions()
    facts = create_facts(dimensions, rng)
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.properties.creator = "KMB BI Portfolio"
    workbook.properties.title = "KMB Synthetic Operations Data"
    workbook.properties.description = "Entirely fictional synthetic source data for a Power BI portfolio demonstration."
    column_orders = {
        "Dim_Date": ["DateKey", "Date", "Year", "Quarter", "MonthNumber", "Month", "YearMonth", "WeekStartDate", "WeekLabel", "DayOfWeekNumber", "DayName", "IsWeekend", "IsCompleteWeek"],
        "Dim_Product": ["ProductKey", "ProductCode", "ProductFamily", "ProductName", "Revision", "StandardCycleMinutes", "StandardUnitValue", "StandardUnitCost"],
        "Dim_WorkCenter": ["WorkCenterKey", "WorkCenterCode", "Line", "Process", "Department"],
        "Dim_Employee": ["EmployeeKey", "EmployeeCode", "EmployeeName", "Role", "Department", "Shift", "HireDate", "IsActive"],
        "Dim_Defect": ["DefectKey", "DefectCode", "DefectName", "Category", "Severity"],
        "Dim_Supplier": ["SupplierKey", "SupplierCode", "SupplierName", "Commodity", "Status"],
        "Fact_Production": ["ProductionEventKey", "DateKey", "ProductKey", "WorkCenterKey", "Shift", "PlannedQty", "ProducedQty", "GoodQty", "ScrapQty", "ReworkQty", "LaborHours", "DowntimeMinutes", "ProductionValue", "PlannedCompletionDateTime", "ActualCompletionDateTime", "DataSource"],
        "Fact_QualityInspection": ["InspectionKey", "ProductionEventKey", "DateKey", "ProductKey", "WorkCenterKey", "DefectKey", "InspectedQty", "DefectQty", "InspectionResult", "DataSource"],
        "Fact_MRB": ["MRBCaseKey", "DateKey", "ProductKey", "WorkCenterKey", "SupplierKey", "DefectKey", "OpenDate", "CloseDate", "Quantity", "MaterialCost", "Disposition", "Status", "DataSource"],
        "Fact_CorrectiveAction": ["ActionKey", "MRBCaseKey", "OpenDate", "DueDate", "CompletionDate", "OwnerDepartment", "Status", "EffectivenessVerified", "DataSource"],
        "Fact_Training": ["TrainingRecordKey", "EmployeeKey", "CourseName", "AssignmentDate", "DueDate", "CompletionDate", "Status", "ScorePct", "DataSource"],
        "Fact_Labor": ["LaborRecordKey", "DateKey", "EmployeeKey", "WorkCenterKey", "RegularHours", "OvertimeHours", "AbsenceHours", "DataSource"],
        "Fact_KpiTargets": ["TargetKey", "KPIName", "PeriodStart", "ScopeType", "ScopeKey", "TargetValue", "Direction", "YellowTolerancePct", "RedTolerancePct", "IsActive", "DataSource"],
    }
    for name, columns in column_orders.items():
        rows_to_sheet(workbook, name, columns, dimensions.get(name, facts.get(name, [])))
    workbook.save(WORKBOOK_PATH)
    sha256 = hashlib.sha256(WORKBOOK_PATH.read_bytes()).hexdigest()
    manifest = {"file": WORKBOOK_PATH.name, "sha256": sha256, "seed": SEED, "data_source": SOURCE_LABEL, "period": {"start": START_DATE.isoformat(), "end": END_DATE.isoformat()}, "row_counts": {name: len(dimensions.get(name, facts.get(name, []))) for name in column_orders}, "generated_at_utc": datetime.now(timezone.utc).replace(microsecond=0).isoformat()}
    MANIFEST_PATH.write_text(json.dumps(manifest, indent=2), encoding="utf-8")
    print(json.dumps(manifest, indent=2))


if __name__ == "__main__":
    generate()
